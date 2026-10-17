from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


def reservations_excel(reservations, filename):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Reservations'
    headers = ['Booking number', 'Full name', 'Phone', 'Email', 'Deaf (account holder)',
               'Tour', 'City', 'Start date', 'End date', 'Days', 'Persons', 'Total',
               'Payment method', 'Payment status', 'Booking status', 'Guest full name',
               'Guest phone', 'Admin note', 'Created on']
    sheet.append(headers)
    for booking in reservations:
        user = booking.user
        profile = getattr(user, 'profile', None)
        sheet.append([
            booking.pk, user.get_full_name() or None, profile.phone or None if profile else None,
            user.email or None, ('Yes' if profile.is_deaf else 'No') if profile else None,
            booking.tour.title, booking.tour.destination.name,
            booking.start_date, booking.end_date, booking.duration_days, booking.num_persons,
            booking.total_price, booking.get_payment_method_display() if booking.payment_method else None,
            booking.get_payment_status_display(), booking.get_status_display(),
            booking.guest_full_name or None if booking.booking_for_other else None,
            booking.guest_phone or None if booking.booking_for_other else None,
            booking.admin_note or None, booking.created_at.replace(tzinfo=None),
        ])
    edge = Side(style='thin', color='CAD8DC')
    widths = [18, 26, 22, 32, 22, 32, 22, 16, 16, 10, 10, 18, 22, 19, 19, 26, 22, 48, 22]
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'  # Keep user text literal, including strings starting with '='.
            cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if cell.row == 1:
                cell.fill = PatternFill('solid', fgColor='193F47')
                cell.font = Font(color='FFFFFF', bold=True)
            else:
                cell.font = Font(name='Calibri', size=11)
                if cell.column in (8, 9):
                    cell.number_format = 'dd mmm yyyy'
                elif cell.column == 12:
                    cell.number_format = '#,##0.00'
                elif cell.column == 19:
                    cell.number_format = 'dd mmm yyyy hh:mm'
    sheet.row_dimensions[1].height = 32
    sheet.freeze_panes = 'D2'
    sheet.auto_filter.ref = sheet.dimensions
    if sheet.max_row > 1:
        report = Table(displayName='Reservations', ref=sheet.dimensions)
        report.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        sheet.add_table(report)
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = '1:1'
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
