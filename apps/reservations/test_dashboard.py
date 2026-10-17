from datetime import date, timedelta
from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase, override_settings
from openpyxl import load_workbook

from apps.core.models import Destination, Tour, Reservation, UserProfile


@override_settings(SECURE_SSL_REDIRECT=False)
class ReservationDashboardTests(TestCase):
    def setUp(self):
        self.staff = User.objects.create_superuser('staff', 'staff@example.com', 'test-password')
        self.user = User.objects.create_user('alex', first_name='Alex', last_name='Martin')
        UserProfile.objects.create(user=self.user, phone='0612345678', country='morocco', postal_code='10000')
        self.tour = Tour.objects.create(title='Rabat', destination=Destination.objects.create(name='Rabat'), price_per_night=100)
        self.rows = []
        for number in range(9):
            start = date.today() + timedelta(days=number * 3 + 10)
            self.rows.append(Reservation.objects.create(user=self.user, tour=self.tour,
                start_date=start, end_date=start+timedelta(days=1), total_price=100))
        self.client.force_login(self.staff)

    def test_search_and_seven_row_pages(self):
        response = self.client.get('/morocco/dashboard/reservations/')
        self.assertEqual(len(response.context['reservations']), 7)
        self.assertEqual(len(self.client.get('/morocco/dashboard/reservations/?page=2').context['reservations']), 2)
        self.assertEqual(self.client.get('/morocco/dashboard/reservations/', {'q': 'Alex Martin'}).context['reservations_count'], 9)
        self.assertEqual(self.client.get('/morocco/dashboard/reservations/', {'q': '0612345678'}).context['reservations_count'], 9)
        self.assertEqual(self.client.get('/morocco/dashboard/reservations/', {'q': 'missing-person'}).context['reservations_count'], 0)
        self.assertNotContains(response, 'id="editFullPackage"')

    def test_excel_is_styled_filtered_and_preserves_blank_cells(self):
        self.user.first_name = '=1+1'
        self.user.last_name = ''
        self.user.save()
        response = self.client.post('/morocco/dashboard/reservations/', {'action': 'download_report', 'reservation_ids': [self.rows[0].pk]})
        self.assertIn('.xlsx', response['Content-Disposition'])
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet['B2'].value, '=1+1')
        self.assertEqual(sheet['B2'].data_type, 's')
        self.assertEqual(sheet['C2'].value, '0612345678')
        self.assertIsNone(sheet['D2'].value)
        self.assertIsNone(sheet['P2'].value)
        self.assertIsNone(sheet['R2'].value)
        self.assertEqual(sheet['A1'].border.bottom.style, 'thin')
        self.assertEqual(sheet.freeze_panes, 'D2')
        response = self.client.post('/morocco/dashboard/reservations/', {'action': 'download_report', 'q': 'missing-person'})
        self.assertEqual(load_workbook(BytesIO(response.content)).active.max_row, 1)
        response = self.client.post('/morocco/dashboard/reservations/', {'action': 'download_report'})
        self.assertEqual(load_workbook(BytesIO(response.content)).active.max_row, 10)

    def test_edit_preserves_removed_options_and_guest_information(self):
        booking = self.rows[0]
        booking.full_package = True
        booking.include_transport = True
        booking.include_hotel = True
        booking.booking_for_other = True
        booking.guest_full_name = 'Other Traveler'
        booking.guest_phone = '0611111111'
        booking.save()
        self.client.post(f'/morocco/dashboard/reservation/{booking.pk}/update/', {
            'start_date': booking.start_date.isoformat(), 'end_date': booking.end_date.isoformat(),
            'num_persons': 2, 'status': 'pending', 'payment_method': 'cash', 'payment_status': 'unpaid', 'admin_note': 'Updated'})
        booking.refresh_from_db()
        self.assertTrue(booking.full_package and booking.include_transport and booking.include_hotel)
        self.assertTrue(booking.booking_for_other)
        self.assertEqual(booking.guest_full_name, 'Other Traveler')
        self.assertEqual(booking.guest_phone, '0611111111')
